import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from difflib import SequenceMatcher
from collections import defaultdict
import os

class PDFProcessor:
    """Process PDF files and convert to Excel with grouped similar lines"""
    
    def __init__(self, similarity_threshold=0.8):
        """
        Initialize PDF processor
        
        Args:
            similarity_threshold: Float between 0-1 for line similarity (default 0.8)
        """
        self.similarity_threshold = similarity_threshold
    
    def extract_text_lines(self, pdf_path):
        """
        Extract text lines from PDF file
        
        Args:
            pdf_path: Path to PDF file
            
        Returns:
            List of text lines
        """
        lines = []
        try:
            with pdfplumber.open(pdf_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        page_lines = [line.strip() for line in text.split('\n') if line.strip()]
                        lines.extend(page_lines)
        except Exception as e:
            raise Exception(f"Error extracting text from PDF: {str(e)}")
        
        return lines
    
    def calculate_similarity(self, str1, str2):
        """
        Calculate similarity ratio between two strings
        
        Args:
            str1: First string
            str2: Second string
            
        Returns:
            Float between 0-1 representing similarity
        """
        return SequenceMatcher(None, str1.lower(), str2.lower()).ratio()
    
    def group_similar_lines(self, lines):
        """
        Group similar lines together
        
        Args:
            lines: List of text lines
            
        Returns:
            List of tuples (group_representative, [similar_lines], count)
        """
        if not lines:
            return []
        
        groups = []
        used_indices = set()
        
        for i, line in enumerate(lines):
            if i in used_indices:
                continue
            
            # Start a new group with this line
            group = [line]
            used_indices.add(i)
            
            # Find similar lines
            for j, other_line in enumerate(lines):
                if j <= i or j in used_indices:
                    continue
                
                similarity = self.calculate_similarity(line, other_line)
                if similarity >= self.similarity_threshold:
                    group.append(other_line)
                    used_indices.add(j)
            
            groups.append((line, group, len(group)))
        
        # Sort by count (descending)
        groups.sort(key=lambda x: x[2], reverse=True)
        
        return groups
    
    def create_excel_from_groups(self, groups, output_path):
        """
        Create Excel file from grouped lines
        
        Args:
            groups: List of tuples (representative, lines, count)
            output_path: Path to save Excel file
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Grouped Lines"
        
        # Headers
        headers = ["Representative Line", "Count", "All Similar Lines"]
        ws.append(headers)
        
        # Style headers
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add data
        for representative, group_lines, count in groups:
            all_lines = " | ".join(group_lines)
            ws.append([representative, count, all_lines])
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 80
        
        # Add styling for data rows
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
            row[1].alignment = Alignment(horizontal="center", vertical="center")
            row[2].alignment = Alignment(wrap_text=True, vertical="top")
        
        # Save workbook
        wb.save(output_path)
    
    def process_pdf(self, pdf_path, output_path):
        """
        Process PDF file and create Excel with grouped lines
        
        Args:
            pdf_path: Path to input PDF file
            output_path: Path to output Excel file
            
        Returns:
            Dictionary with processing stats
        """
        try:
            # Extract lines from PDF
            lines = self.extract_text_lines(pdf_path)
            
            if not lines:
                raise Exception("No text found in PDF")
            
            # Group similar lines
            groups = self.group_similar_lines(lines)
            
            # Create Excel file
            self.create_excel_from_groups(groups, output_path)
            
            return {
                'success': True,
                'total_lines': len(lines),
                'unique_groups': len(groups),
                'output_path': output_path
            }
        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }
